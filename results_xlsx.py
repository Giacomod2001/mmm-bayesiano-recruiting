"""
Output unico in Excel: un solo workbook con tutti i risultati finali.

Ogni stage della pipeline scrive il proprio foglio nello stesso file
`pipeline/data/output/risultati.xlsx` invece di un CSV separato:

    Canali           allocator stage 1 (quarter per canale)
    Settimane        allocator: spaccato settimanale
    Mesi             allocator: rollup mensile
    Campagne         allocator stage 2 (riparto per campagna)
    Recovery         parameter recovery (stima vs verità)

I CSV "grezzi" del generatore (meta_ads, google_ads, ...) NON passano da
qui: servono apposta a testare l'ingestion e restano nel formato messy
degli export reali. Anche i canonici (media/outcome/seasonality.csv)
restano CSV perché vengono riletti dagli stage successivi.

Il file viene aggiornato in-place: scrivere un foglio già presente lo
rimpiazza, gli altri restano. Così non importa l'ordine con cui si
lanciano i tre script — il workbook si compone via via.
"""
from __future__ import annotations

import os

import pandas as pd

# pipeline/data/output accanto a questo file (repo root)
ROOT = os.path.dirname(os.path.abspath(__file__))
# Anche il percorso del workbook è sovrascrivibile da ambiente, per la stessa
# ragione di `config.OUTPUT_DIR`: il collaudo scrive un Excel di prova senza
# sovrascrivere quello di lavoro.
WORKBOOK = os.environ.get(
    "MMM_WORKBOOK",
    os.path.join(ROOT, "pipeline", "data", "output", "risultati.xlsx"))

MONEY = "#,##0 €"


def write_sheet(sheet: str, df: pd.DataFrame,
                formats: dict[str, str] | None = None,
                path: str = WORKBOOK) -> str:
    """Scrive (o rimpiazza) un foglio nel workbook unico, con larghezze
    colonna sensate e number_format per le colonne indicate.

    `formats`: mappa colonna -> formato Excel (es. {"spend": MONEY}).
    Le date in colonna "week" vengono normalizzate a data pura.
    """
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    out = df.copy()
    if "week" in out.columns:
        out["week"] = pd.to_datetime(out["week"]).dt.date

    if os.path.exists(path):
        kw = dict(engine="openpyxl", mode="a", if_sheet_exists="replace")
    else:
        kw = dict(engine="openpyxl", mode="w")

    with pd.ExcelWriter(path, **kw) as xl:
        out.to_excel(xl, sheet_name=sheet, index=False)
        ws = xl.sheets[sheet]
        for j, col in enumerate(out.columns, start=1):
            letter = ws.cell(1, j).column_letter
            ws.column_dimensions[letter].width = max(len(str(col)) + 2, 14)
            fmt = (formats or {}).get(col)
            if fmt:
                for i in range(2, len(out) + 2):
                    ws.cell(i, j).number_format = fmt
    return path


def add_images(sheet: str, image_paths: list[str], path: str = WORKBOOK) -> str:
    """Inserisce immagini PNG in un foglio del workbook (impilate in verticale),
    creando o rimpiazzando il foglio. Richiede Pillow. Da chiamare DOPO che il
    workbook esiste (write_sheet)."""
    from openpyxl import load_workbook
    from openpyxl.drawing.image import Image as XLImage

    wb = load_workbook(path)
    if sheet in wb.sheetnames:
        del wb[sheet]
    ws = wb.create_sheet(sheet)
    row = 1
    for p in image_paths:
        if os.path.exists(p):
            ws.add_image(XLImage(p), f"A{row}")
            row += 26                      # spazio verticale per l'immagine
    wb.save(path)
    return path
