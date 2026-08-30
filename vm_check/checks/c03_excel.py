"""
Controllo 3 — Excel: lettura degli export .xlsx e scrittura del workbook
dei risultati.

Sulla VM l'Excel è il punto in cui le cose si rompono più spesso (openpyxl
mancante, motore sbagliato, formati numerici persi). Qui si verifica:
  1. lettura dei due export .xlsx sporchi (LinkedIn, richieste clienti);
  2. scrittura di più fogli nello stesso workbook con `results_xlsx`;
  3. aggiornamento in-place di un foglio senza perdere gli altri;
  4. tenuta del formato valuta e dei valori dopo la rilettura.
"""
from __future__ import annotations

import os

from ._util import Result, Timer, fail, ok, skip, verdict

TITOLO = "Excel"

XLSX_ATTESI = ("linkedin_campaigns.xlsx", "richieste_clienti.xlsx")


def run(ctx: dict) -> list[Result]:
    import pandas as pd
    from openpyxl import load_workbook

    import results_xlsx as rx

    res: list[Result] = []
    out_dir = ctx["out_dir"]
    raw_dir = ctx.get("raw_dir")

    # --- 1. lettura degli export .xlsx --------------------------------------
    if not raw_dir:
        res.append(skip("Lettura export .xlsx",
                        "il controllo ingestion non ha prodotto dati"))
    else:
        for nome in XLSX_ATTESI:
            percorso = os.path.join(raw_dir, nome)
            if not os.path.exists(percorso):
                res.append(fail(f"Lettura {nome}", "file non trovato"))
                continue
            with Timer() as t:
                try:
                    fogli = pd.read_excel(percorso, sheet_name=None,
                                          engine="openpyxl")
                except Exception as e:  # noqa: BLE001
                    res.append(fail(f"Lettura {nome}",
                                    f"{type(e).__name__}: {e}", t.elapsed))
                    continue
            righe = sum(len(d) for d in fogli.values())
            res.append(verdict(
                f"Lettura {nome}", righe > 0,
                f"{len(fogli)} foglio/i, {righe:,} righe",
                "file letto ma vuoto", seconds=t.elapsed))

    # --- 2. scrittura multi-foglio ------------------------------------------
    wb_path = os.path.join(out_dir, "prova_workbook.xlsx")
    if os.path.exists(wb_path):
        try:
            os.remove(wb_path)
        except OSError:                # mount senza permesso di cancellazione:
            wb_path = os.path.join(    # si scrive su un nome nuovo
                out_dir, f"prova_workbook_{os.getpid()}.xlsx")

    media = ctx.get("facts", {}).get("media")
    if media is not None and len(media):
        canali = (media.groupby("channel", as_index=False)["spend"].sum()
                  .rename(columns={"spend": "Spesa"}))
        settimane = (media.groupby("week", as_index=False)["spend"].sum()
                     .rename(columns={"spend": "Spesa"}).head(20))
    else:                              # dati finti: il controllo resta valido
        canali = pd.DataFrame({"channel": ["google", "meta", "linkedin"],
                               "Spesa": [1000.0, 750.5, 250.25]})
        settimane = pd.DataFrame(
            {"week": pd.date_range("2024-01-01", periods=5, freq="W-MON"),
             "Spesa": [100.0, 110.0, 120.0, 130.0, 140.0]})

    with Timer() as t:
        try:
            rx.write_sheet("Canali", canali, {"Spesa": rx.MONEY}, path=wb_path)
            rx.write_sheet("Settimane", settimane, {"Spesa": rx.MONEY},
                           path=wb_path)
        except Exception as e:  # noqa: BLE001
            res.append(fail("Scrittura workbook multi-foglio",
                            f"{type(e).__name__}: {e}", t.elapsed))
            return res
    res.append(ok("Scrittura workbook multi-foglio",
                  f"2 fogli in {os.path.basename(wb_path)}", t.elapsed))

    # --- 3. aggiornamento in-place ------------------------------------------
    canali_agg = canali.copy()
    canali_agg["Spesa"] = canali_agg["Spesa"] * 2
    try:
        rx.write_sheet("Canali", canali_agg, {"Spesa": rx.MONEY}, path=wb_path)
    except Exception as e:  # noqa: BLE001
        res.append(fail("Aggiornamento foglio esistente",
                        f"{type(e).__name__}: {e}"))
    else:
        wb = load_workbook(wb_path)
        conserva = "Settimane" in wb.sheetnames and "Canali" in wb.sheetnames
        res.append(verdict(
            "Aggiornamento foglio esistente", conserva,
            "il foglio riscritto rimpiazza sé stesso e gli altri restano",
            f"fogli presenti dopo la riscrittura: {wb.sheetnames}"))
        wb.close()

    # --- 4. rilettura: valori e formati -------------------------------------
    letto = pd.read_excel(wb_path, sheet_name="Canali", engine="openpyxl")
    atteso = float(canali_agg["Spesa"].sum())
    trovato = float(letto["Spesa"].sum())
    res.append(verdict(
        "Valori conservati dopo la rilettura",
        abs(trovato - atteso) < max(1e-6, abs(atteso) * 1e-9),
        f"totale {trovato:,.2f}",
        f"scritto {atteso:,.2f} ma riletto {trovato:,.2f}"))

    wb = load_workbook(wb_path)
    ws = wb["Canali"]
    col = [c.value for c in ws[1]].index("Spesa") + 1
    formato = ws.cell(2, col).number_format
    wb.close()
    res.append(verdict("Formato valuta conservato", formato == rx.MONEY,
                       f"number_format = {formato!r}",
                       f"atteso {rx.MONEY!r}, trovato {formato!r}"))

    ctx["workbook_prova"] = wb_path
    return res
